import csv, glob, re, os, collections
from openpyxl import Workbook, load_workbook

csv_f = glob.glob('*1208.csv')[0]
xlsx_f = 'AXA_IM_disclaimers_with_ISIN.xlsx'
bad = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
isin_re = re.compile(r'^[A-Z]{2}[A-Z0-9]{9}[0-9]$')
num_cols = ('disclaimerId', 'linkedEntityId')

with open(csv_f, newline='', encoding='utf-8-sig') as fh:
    r = csv.reader(fh)
    header = next(r)
    rows = list(r)
ix = {h: i for i, h in enumerate(header)}

widths = {'disclaimerId': 14, 'disclaimerType': 28, 'linkedEntityId': 15, 'isin': 15, 'culture': 10,
          'disclaimer': 80, 'lastModified': 20, 'created': 20, 'status': 18, 'lastPublished': 20,
          'lastAction': 18, 'lastModifiedBy': 32, 'createdBy': 32, 'lastPublishedBy': 32}
wb = Workbook(write_only=True)
ws = wb.create_sheet('Disclaimers')
ws.freeze_panes = 'A2'
for i, h in enumerate(header):
    ws.column_dimensions[chr(65 + i)].width = widths.get(h, 18)
ws.append(header)
for x in rows:
    out = []
    for h, c in zip(header, x):
        c = bad.sub('', c)
        out.append(int(c) if h in num_cols and c.isdigit() else c)
    ws.append(out)
wb.save(xlsx_f)
print("Spreadsheet saved:", xlsx_f)
print()

print("CSV file:", csv_f)
print("Columns:", header)
print("Data rows:", len(rows), "(expected 104590)")
print("Wrong column count (should be 0):", sum(1 for x in rows if len(x) != len(header)))
print("Status:", dict(collections.Counter(x[ix['status']] for x in rows)))
print("Status/lastAction:", dict(collections.Counter((x[ix['status']], x[ix['lastAction']]) for x in rows)))
blank = [x for x in rows if not x[ix['isin']]]
print("Blank ISIN rows:", len(blank), "(expected about 3238)")
print("  blank ISIN but has linkedEntityId (should be 0):", sum(1 for x in blank if x[ix['linkedEntityId']]))
print("  ISIN but no linkedEntityId (should be 0):", sum(1 for x in rows if x[ix['isin']] and not x[ix['linkedEntityId']]))
odd = {x[ix['isin']] for x in rows if x[ix['isin']] and not isin_re.match(x[ix['isin']])}
print("Non-standard ISINs:", len(odd), list(odd)[:5])
print("Garbled characters (should be 0):", sum(1 for x in rows if any(s in x[ix['disclaimer']] for s in ('Ã', 'â€', '√'))))
print("Blank disclaimer text (should be 0):", sum(1 for x in rows if not x[ix['disclaimer']].strip()))
print("Blank disclaimerType (should be 0):", sum(1 for x in rows if not x[ix['disclaimerType']]))
print("Blank status (should be 0):", sum(1 for x in rows if not x[ix['status']]))
keys = collections.Counter((x[ix['disclaimerId']], x[ix['culture']]) for x in rows)
dups = [k for k, v in keys.items() if v > 1]
same = diff = 0
for k in dups:
    t = {x[ix['disclaimer']] for x in rows if (x[ix['disclaimerId']], x[ix['culture']]) == k}
    if len(t) == 1: same += 1
    else: diff += 1
print("Duplicate disclaimerId+culture pairs:", len(dups), "| identical text:", same, "| different text:", diff)
print("  examples:", dups[:5])
print("Distinct disclaimerIds:", len({x[ix['disclaimerId']] for x in rows}), "(expected 7865)")
print("Distinct ISINs:", len({x[ix['isin']] for x in rows if x[ix['isin']]}), "(expected 2431)")
print("Distinct types:", len({x[ix['disclaimerType']] for x in rows}), "(expected 203)")

ws2 = load_workbook(xlsx_f, read_only=True).active
n = mism = 0
for i, xr in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0: continue
    n += 1
    xr = [('' if v is None else str(v)) for v in xr]
    if xr != [bad.sub('', c) for c in rows[i - 1]]: mism += 1
print("XLSX data rows:", n, "| equals CSV:", n == len(rows), "| differing rows (should be 0):", mism)
