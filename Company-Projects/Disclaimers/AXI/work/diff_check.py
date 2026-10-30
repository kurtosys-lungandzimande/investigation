import csv, re, collections
from openpyxl import load_workbook
bad = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
with open('AXA_IM_disclaimers_with_ISIN.csv', newline='', encoding='utf-8-sig') as fh:
    r = csv.reader(fh); header = next(r); rows = list(r)
ws = load_workbook('AXA_IM_disclaimers_with_ISIN.xlsx', read_only=True).active
per = collections.Counter(); shown = 0; short = 0
for i, xr in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    xr = [('' if v is None else str(v)) for v in xr]
    if len(xr) < len(header):
        short += 1
        xr += [''] * (len(header) - len(xr))
    cr = [bad.sub('', c) for c in rows[i-1]]
    for j, h in enumerate(header):
        if xr[j] != cr[j]:
            per[h] += 1
            if shown < 5:
                shown += 1
                print('Row', i+1, h, 'xlsx=', repr(xr[j][:80]), 'csv=', repr(cr[j][:80]))
print('Rows shorter than header in xlsx:', short)
print('Differences by column:', dict(per))
