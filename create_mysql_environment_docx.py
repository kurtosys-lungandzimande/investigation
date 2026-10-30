#!/usr/bin/env python3
"""Generate Word document for MySQL/RDS environment session notes."""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUTPUT = "/Users/lunga.ndzimande/Documents/MySQL-RDS-Environment-Session-Notes.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], val)
    doc.add_paragraph()


def main():
    doc = Document()

    title = doc.add_heading("MySQL / RDS Environment — Session Notes", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("High-level overview (anonymized)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    doc.add_paragraph(
        "Separate document from AWS environments overview and change management notes. "
        "No company, product, or person names are included."
    )

    # 1
    doc.add_heading("1. What this session covers", level=1)
    for item in [
        "High-level overview of the MySQL (RDS) environment",
        "Where environments sit (regions, tiers)",
        "Normal procedures and how the platform works",
        "Backups and related topics (mentioned — covered later in the full presentation)",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # 2
    doc.add_heading("2. What the platform does (business view)", level=1)
    doc.add_paragraph(
        "The MySQL environment supports a digital experience / website management platform. "
        "It is a core part of a larger application stack."
    )
    for item in [
        "Lets design hosts and clients manage their websites.",
        "Hosts WordPress-based sites.",
        "Clients build sites in the application; site data is stored in the database.",
        "Sites serve dynamic, data-driven content to end users.",
        "Sites can be published live for visitors.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("How it connects to the wider stack", level=2)
    for item in [
        "Connects to a separate backend WordPress application.",
        "Pulls financial and marketing data for clients to present on their websites.",
        "End users see content via the website; they do not touch the database directly.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # 3
    doc.add_heading("3. Who uses it & design goals", level=1)
    add_table(
        doc,
        ["Audience", "Need"],
        [
            ["Marketing teams", "Manage website content"],
            ["Developers", "Maintain backend; improvements and enhancements"],
            ["Clients / end users", "Reliable, secure, data-driven websites"],
        ],
    )
    doc.add_paragraph("Design intent:", style="List Bullet")
    doc.add_paragraph("Low maintenance — focus on performance, not constant firefighting", style="List Bullet 2")
    doc.add_paragraph("Secure and reliable for client data", style="List Bullet 2")

    # 4
    doc.add_heading("4. Database & tenancy model", level=1)
    doc.add_paragraph("RDS is multi-tenant.", style="List Bullet")
    doc.add_paragraph(
        "Multiple clients/environments share the managed database service model; "
        "logical separation is enforced per environment tier and region."
    )

    # 5
    doc.add_heading("5. Environment tiers (dev → staging → production)", level=1)
    add_table(
        doc,
        ["Tier", "Purpose"],
        [
            ["Development", "Teams test changes, content updates, upgrades, enhancements"],
            ["Staging", "Changes promoted after dev; review and confirm before prod"],
            ["Production", "Live sites for real end users"],
        ],
    )
    doc.add_paragraph(
        "All tiers are isolated from each other to prevent cross-contamination and keep stability."
    )

    # 6
    doc.add_heading("6. Regions & data sovereignty", level=1)
    doc.add_paragraph("Production is split across three regions:", style="List Bullet")
    doc.add_paragraph("UK", style="List Bullet 2")
    doc.add_paragraph("US", style="List Bullet 2")
    doc.add_paragraph("EU", style="List Bullet 2")

    doc.add_heading("Data sovereignty (critical rule)", level=2)
    for item in [
        "Data does NOT move between regions (e.g. UK data stays in UK only).",
        "Clients are aware of and expect this — data stays in the region where it is stored.",
        "If a client is in UK, their data remains UK-only.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Backups (Q&A from session)", level=2)
    doc.add_paragraph(
        "Backups exist and align with client/regional requirements; details were to be "
        "covered on a later slide in the same presentation. Confirm backup policy per region "
        "with the team or runbook."
    )

    # 7
    doc.add_heading("7. Hosting on AWS", level=1)
    for item in [
        "Platform is hosted on AWS.",
        "Multiple regions for data sovereignty (UK, US, EU).",
        "Within each region: multiple availability zones (AZs).",
        "AZs give high availability — if one zone fails, data/services in another zone support continuity.",
        "Production uses multi-AZ for redundancy, availability, and performance.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # 8
    doc.add_heading("8. How people access the database", level=1)
    add_table(
        doc,
        ["Who", "How they access data"],
        [
            ["Clients / public", "Via the WordPress website only — NOT direct DB access"],
            ["Internal team", "Via a jump box → then RDS (internal use only)"],
        ],
    )

    doc.add_heading("Jump box (internal access path)", level=2)
    for item in [
        "Jump box connects you through to RDS.",
        "Still requires RDS username and password — not an open connection.",
        "Clients have no access to jump box or direct database connections.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Flow:", style="List Bullet")
    p = doc.add_paragraph(style="List Bullet 2")
    p.add_run("Client → WordPress site → application layer → database")

    p2 = doc.add_paragraph(style="List Bullet 2")
    p2.add_run("Engineer → Jump box (credentials) → RDS (username/password)")

    # 9
    doc.add_heading("9. Operations & reliability", level=1)
    for item in [
        "Environment is scalable.",
        "Monitored for high uptime.",
        "Changes follow dev → staging → production promotion.",
        "Production changes use the formal change-management process (see separate notes doc).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # 10
    doc.add_heading("10. Architecture diagram (mental model)", level=1)
    diagram = """  ┌─────────────┐     ┌──────────────────┐     ┌─────────────┐
  │   Client    │────▶│  WordPress site  │────▶│  RDS MySQL  │
  │  (browser)  │     │  (no direct DB)  │     │  (multi-AZ) │
  └─────────────┘     └──────────────────┘     └──────▲──────┘
                                                      │
  ┌─────────────┐     ┌──────────────────┐           │
  │  Engineer   │────▶│    Jump box      │───────────┘
  │  (internal) │     │  + DB credentials│
  └─────────────┘     └──────────────────┘

  Regions: UK | US | EU  —  data never crosses regions
  Tiers:   Dev → Staging → Production  —  isolated"""
    p = doc.add_paragraph()
    run = p.add_run(diagram)
    run.font.name = "Courier New"
    run.font.size = Pt(9)

    # 11
    doc.add_heading("11. Session Q&A captured", level=1)
    doc.add_paragraph("Question: If data stays in UK for sovereignty, what about backups?")
    doc.add_paragraph(
        "Answer (from session): Backups are in place and covered later in the deck; "
        "they should respect the same regional/sovereignty rules — verify with team docs."
    )

    # 12
    doc.add_heading("12. Terms (plain English)", level=1)
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ["RDS", "AWS managed relational database (MySQL here)"],
            ["Multi-tenant", "Shared infrastructure; separated logically per client/env"],
            ["Data sovereignty", "Data stays in the region where it was provisioned"],
            ["Availability zone (AZ)", "Separate datacenter area within a region for failover"],
            ["Jump box", "Bastion host — secure hop for internal DB access only"],
            ["Staging", "Pre-production environment for final validation"],
        ],
    )

    # Related docs
    doc.add_heading("Related notes (other documents)", level=1)
    add_table(
        doc,
        ["Topic", "Document file"],
        [
            ["All AWS DB environments (overview)", "AWS-Database-Environments-Session-Notes.docx"],
            ["Change management & tickets", "Change-Management-Session-Notes.docx"],
            ["This session (MySQL / RDS detail)", "MySQL-RDS-Environment-Session-Notes.docx"],
        ],
    )

    doc.add_paragraph()
    footer = doc.add_paragraph(
        "Generated for personal reference. Anonymized — no sensitive identifiers. "
        "Transcript may continue — add Part 2 when more content is available."
    )
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
