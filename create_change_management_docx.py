#!/usr/bin/env python3
"""Generate Word document for change management session notes."""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUTPUT = "/Users/lunga.ndzimande/Documents/Change-Management-Session-Notes.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], val)
    doc.add_paragraph()


def main():
    doc = Document()

    title = doc.add_heading("Change Management & Change Tickets — Session Notes", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Knowledge share (anonymized)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    doc.add_paragraph(
        "Separate document from the AWS database environments overview. "
        "No company, product, or person names are included."
    )

    # Section 1
    doc.add_heading("1. What is change management?", level=1)
    doc.add_paragraph(
        "Any work that changes production databases must go through a formal change process."
    )
    for item in [
        "The organisation knows exactly what will change and where.",
        "You get official permission to perform the change.",
        "You cannot make ad-hoc changes directly in production — production affects clients.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # Section 2
    doc.add_heading("2. Production regions & timing", level=1)
    doc.add_paragraph("Production is split across three regions:", style="List Bullet")
    doc.add_paragraph("EU", style="List Bullet 2")
    doc.add_paragraph("UK", style="List Bullet 2")
    doc.add_paragraph("US", style="List Bullet 2")

    doc.add_heading("Change windows", level=2)
    for item in [
        "Changes are only allowed before or after normal business working hours.",
        "Goal: avoid disrupting clients.",
        "If something goes wrong, rollback must be possible within the same window.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph(
        "For patching: there are usually defined maintenance windows, and clients may "
        "be notified by email before work begins."
    )

    # Section 3
    doc.add_heading("3. Where changes are logged (the change ticket / CC)", level=1)
    doc.add_paragraph(
        "Changes are logged in a tracking system under your team’s column/queue "
        "(e.g. database engineering). Each production change gets its own change ticket (CC)."
    )
    doc.add_paragraph(
        "Example type of work: patching production database instances."
    )

    # Section 4
    doc.add_heading("4. What to put in the change ticket", level=1)
    doc.add_paragraph("Be as descriptive as possible. The ticket should stand alone so someone else can understand and execute it.")

    doc.add_heading("Required / expected content", level=2)
    add_table(
        doc,
        ["Field / section", "What to include"],
        [
            ["Description", "Everything you will do during the change"],
            ["Scripts", "Attach or link any SQL/scripts that will run"],
            ["Environments", "Which regions/instances will be touched (EU, UK, US)"],
            ["Deployment window", "Exact date and time for each environment"],
            ["High-level steps", "Step-by-step plan of the work"],
            ["Reason for change", "Why this change is needed"],
            ["Risk", "Known risks; anything sensitive that could be affected"],
            ["Acceptance criteria", "What “done” looks like; what you expect after the change"],
            ["Deployment team", "Who deploys; who joins the call for backup"],
            ["Testing resources", "Who validates prod after the change"],
            ["Standby resources", "Extra people on call if needed"],
        ],
    )

    doc.add_heading("Example scheduling pattern", level=2)
    doc.add_paragraph(
        "One change might cover UK, US, and EU on a weekend date, scheduled before "
        "working hours. Each region can have its own time slot inside the ticket."
    )

    # Section 5
    doc.add_heading("5. Roles on a production change", level=1)
    add_table(
        doc,
        ["Role", "Purpose"],
        [
            ["Deployer", "One person from the team who executes the change"],
            ["Meeting backup", "Another team member on the call if help is needed"],
            ["Testing resource", "Confirms clients/instances are OK after the change"],
            ["Standby", "Available if extra support is required"],
            ["Peer reviewer", "Reviews the ticket before submission — confirms it is complete and correct"],
        ],
    )

    doc.add_paragraph(
        "Example (data deletion): deployer runs the change; tester checks only the "
        "intended data was removed and the client instance is healthy."
    )

    # Section 6
    doc.add_heading("6. Approval workflow (ticket states)", level=1)
    doc.add_paragraph("Typical flow from creation to completion:")

    states = [
        ("1. Create", "Draft the CC with all fields filled in."),
        ("2. Peer review", "A teammate reviews and confirms the ticket is ready."),
        ("3. Waiting approval", "Submit for approval; ask the right approver."),
        ("4. More information", "If rejected or incomplete — update and resubmit."),
        ("5. Approved", "Wait until the scheduled execution date/time."),
        ("6. Execute", "Perform the change during the approved window."),
        ("7. Deployed / completed", "After change is done, tested, and signed off."),
    ]
    for label, desc in states:
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{label}: ").bold = True
        p.add_run(desc)

    doc.add_paragraph(
        "Important: you may only execute after the ticket is approved and the scheduled time has arrived."
    )

    doc.add_heading("Ticket status reference", level=2)
    add_table(
        doc,
        ["Status", "When to use it"],
        [
            ["To-do", "CC created; work not started or still being prepared"],
            ["Waiting approval", "CC is complete; submitted for approver"],
            ["More information", "Approver needs fixes or extra detail — update and resubmit"],
            ["Approved", "Ready to run on the scheduled date/time"],
            ["In progress / executing", "Posted start notice; change work underway (see live comms)"],
            ["Deployed", "Change done, tested, and signed off"],
        ],
    )

    doc.add_paragraph(
        "Note: The recording ends while describing what to do once approved and the start "
        "notice is posted — typically you move the ticket to an in-progress state when execution begins."
    )

    # Section 7
    doc.add_heading("7. Live communication during the change", level=1)
    doc.add_paragraph(
        "Use the team’s change-notifications channel (e.g. chat app) for every production change."
    )

    doc.add_heading("When starting", level=2)
    for item in [
        "Post in the change-notifications channel for every production change ticket.",
        "Post at the scheduled time — say you are starting now.",
        "Include the ticket link.",
        "Include the ticket ID and title (copy the title from the ticket).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Message templates (copy and adapt):", style="List Bullet")
    p1 = doc.add_paragraph(style="List Bullet 2")
    p1.add_run("Short: ").bold = True
    p1.add_run("“Starting [ticket link] — [brief description]”")
    p2 = doc.add_paragraph(style="List Bullet 2")
    p2.add_run("Detailed: ").bold = True
    p2.add_run("“Starting CC-XXXXX — [paste ticket title here]”")

    doc.add_heading("When finished", level=2)
    doc.add_paragraph(
        "After the change is done, tested, and verified, move the ticket to deployed/completed state.",
        style="List Bullet",
    )

    # Section 8 - session flow
    doc.add_heading("8. Session flow (what was covered)", level=1)
    add_table(
        doc,
        ["Topic", "Key takeaway"],
        [
            ["Why changes need approval", "Production DB changes can affect clients — no ad-hoc changes"],
            ["Regions", "EU, UK, US — schedule each if needed"],
            ["When", "Before/after business hours; rollback must fit in the same window"],
            ["Where tickets live", "Team column/queue in the change tracking tool"],
            ["Ticket content", "Description, scripts, environments, times, steps, risks, acceptance criteria"],
            ["People", "Deployer + backup on call + testers + standby + peer reviewer"],
            ["Approval", "Peer review → submit → approved → then execute on schedule"],
            ["Comms", "Post start (and finish) in change-notifications channel"],
            ["Close-out", "Test → move ticket to Deployed"],
        ],
    )

    # Section 9
    doc.add_heading("9. End-to-end checklist", level=1)
    checklist = [
        "Draft change ticket with full description, scripts, environments, times, risks, acceptance criteria",
        "Assign deployer, backup, testers, standby",
        "Peer review completed",
        "Submit for approval → approved",
        "Notify clients if required (e.g. patching)",
        "Post “starting” message in change-notifications channel at scheduled time",
        "Execute within the change window",
        "Test with testing resources",
        "Rollback if needed (still within window)",
        "Post completion / move ticket to deployed",
    ]
    for item in checklist:
        doc.add_paragraph(item, style="List Bullet")

    # Section 10
    doc.add_heading("10. Quick reference diagram", level=1)
    diagram = """  CREATE CC  →  PEER REVIEW  →  SUBMIT  →  APPROVED?
                    │                              │
                    │                         NO → MORE INFO → fix → resubmit
                    │                              │
                    │                         YES → wait for window
                    ▼
            POST "STARTING" in notifications channel
                    ▼
            EXECUTE (deployer) + BACKUP on call
                    ▼
            TEST (testing resources)
                    ▼
            DONE → move to DEPLOYED + notify"""
    p = doc.add_paragraph()
    run = p.add_run(diagram)
    run.font.name = "Courier New"
    run.font.size = Pt(9)

    # Glossary
    doc.add_heading("11. Terms (plain English)", level=1)
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ["CC / change ticket", "Formal record and approval for a production change"],
            ["Change window", "Approved time slot (usually outside business hours)"],
            ["Deployer", "Person who runs the change"],
            ["Peer review", "Teammate checks the ticket before approval"],
            ["Acceptance criteria", "How you know the change succeeded"],
            ["Rollback", "Revert if something fails — within the same window"],
        ],
    )

    doc.add_paragraph()
    footer = doc.add_paragraph(
        "Generated for personal reference. Anonymized — no sensitive identifiers."
    )
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
